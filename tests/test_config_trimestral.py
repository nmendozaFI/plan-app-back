"""V21 / F3a + V22 (Cambio A): tests for the config-trimestral router.

Endpoints under test:
  - GET /{trimestre}/empresas-ep              (V21): EP+activa filter.
  - GET /{trimestre}/empresas-permite-extras  (V22): permiteExtras+activa filter.
  - GET /{trimestre}                          (V22): row exposes permite_extras.
  - GET /{trimestre}/resumen                  (V22): counter `permite_extras`.
  - PUT /{trimestre}/batch                    (V22): updates permite_extras.
  - PUT /{trimestre}/{empresa_id}             (V22): updates permite_extras.
  - POST /{trimestre}/inicializar             (V22): heredera de aceptaExtras y
                                                     clona permiteExtras.
"""

import pytest
from sqlalchemy import text

from .conftest import TEST_TRIMESTRE, TEST_EMPRESA_PREFIX


# ── Helpers ────────────────────────────────────────────────────


async def _create_empresa(
    db,
    nombre: str,
    activa: bool = True,
    *,
    acepta_extras: bool = False,
) -> int:
    """Insert (or upsert) a test empresa. Returns its id.

    V22 (Cambio A): `acepta_extras` toggles the empresa-level baseline that
    `inicializar_configs` mode-defaults inherits into CT.permiteExtras.
    """
    res = await db.execute(
        text(
            'INSERT INTO empresa (nombre, tipo, activa, "aceptaExtras", "updatedAt") '
            "VALUES (:n, 'AMBAS', :activa, :ae, NOW()) "
            'ON CONFLICT (nombre) DO UPDATE '
            'SET activa = EXCLUDED.activa, "aceptaExtras" = EXCLUDED."aceptaExtras" '
            "RETURNING id"
        ),
        {"n": nombre, "activa": activa, "ae": acepta_extras},
    )
    eid = res.scalar()
    await db.commit()
    return eid


async def _set_config_trimestral(
    db,
    empresa_id: int,
    escuela_propia: bool,
    *,
    permite_extras: bool = False,
    trimestre: str = TEST_TRIMESTRE,
):
    """Create or update configTrimestral for (empresa, trimestre).

    V22 (Cambio A): added `permite_extras` kwarg (independent of escuela_propia).
    """
    await db.execute(
        text(
            'INSERT INTO "configTrimestral" '
            '("empresaId", trimestre, "tipoParticipacion", "escuelaPropia", '
            '"permiteExtras", "disponibilidadDias", "updatedAt") '
            "VALUES (:eid, :tri, 'AMBAS', :ep, :pe, 'L,M,X,J,V', NOW()) "
            'ON CONFLICT ("empresaId", trimestre) DO UPDATE '
            'SET "escuelaPropia" = EXCLUDED."escuelaPropia", '
            '    "permiteExtras" = EXCLUDED."permiteExtras"'
        ),
        {
            "eid": empresa_id,
            "tri": trimestre,
            "ep": escuela_propia,
            "pe": permite_extras,
        },
    )
    await db.commit()


async def _set_puede_ser_ep(db, empresa_id: int, value: bool = True):
    """V25 Cambio C (Capa 3): flag estructural en ficha de empresa.
    Reemplaza el setup viejo `_set_config_trimestral(..., escuela_propia=True)`
    para los tests del endpoint `/empresas-ep` (que ya no pasan por CT)."""
    await db.execute(
        text('UPDATE empresa SET "puedeSerEP" = :v WHERE id = :id'),
        {"id": empresa_id, "v": value},
    )
    await db.commit()


async def _set_puede_ser_doble(db, empresa_id: int, value: bool = True):
    """V25 Cambio C (Capa 3): flag estructural Doble en ficha de empresa.
    Usado por los tests nuevos del endpoint `/empresas-doble`."""
    await db.execute(
        text('UPDATE empresa SET "puedeSerDoble" = :v WHERE id = :id'),
        {"id": empresa_id, "v": value},
    )
    await db.commit()


# ── Tests ──────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_listar_empresas_ep_devuelve_solo_ep_activas(client, db_session):
    """V25 Cambio C (Capa 3): el gate es `empresa.puedeSerEP`, no `CT.escuelaPropia`.
    3 empresas: EP+activa, EP+inactiva, noEP+activa → solo la primera."""
    ep_activa_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}EP_F3A_OK", activa=True
    )
    ep_inactiva_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}EP_F3A_INACTIVA", activa=False
    )
    no_ep_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}EP_F3A_NOEP", activa=True
    )
    await _set_puede_ser_ep(db_session, ep_activa_id)
    await _set_puede_ser_ep(db_session, ep_inactiva_id)
    # no_ep_id keeps default puedeSerEP=false.

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-ep")
    assert resp.status_code == 200, resp.text
    data = resp.json()

    assert data["trimestre"] == TEST_TRIMESTRE
    returned_ids = [e["id"] for e in data["empresas"]]

    # Only the EP+activa empresa from this test should appear among the test ones.
    assert ep_activa_id in returned_ids
    assert ep_inactiva_id not in returned_ids
    assert no_ep_id not in returned_ids

    # Sanity on shape of the row we know is there.
    row = next(e for e in data["empresas"] if e["id"] == ep_activa_id)
    assert row["nombre"] == f"{TEST_EMPRESA_PREFIX}EP_F3A_OK"
    assert row["tipo"] == "AMBAS"
    assert row["activa"] is True

    # `total` must match the array length.
    assert data["total"] == len(data["empresas"])


@pytest.mark.asyncio
async def test_listar_empresas_ep_orden_alfabetico(client, db_session):
    """Three EP empresas inserted out of order → response sorted A, B, C."""
    # Names crafted so they sort A < B < C among themselves.
    name_b = f"{TEST_EMPRESA_PREFIX}ORDEN_B_F3A"
    name_a = f"{TEST_EMPRESA_PREFIX}ORDEN_A_F3A"
    name_c = f"{TEST_EMPRESA_PREFIX}ORDEN_C_F3A"

    # Insert in non-alphabetical order (B, A, C) to ensure ordering comes from
    # the SQL, not from insertion order.
    b_id = await _create_empresa(db_session, name_b)
    a_id = await _create_empresa(db_session, name_a)
    c_id = await _create_empresa(db_session, name_c)

    for eid in (b_id, a_id, c_id):
        await _set_puede_ser_ep(db_session, eid)

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-ep")
    assert resp.status_code == 200, resp.text
    data = resp.json()

    test_rows = [e for e in data["empresas"] if e["nombre"].startswith(
        f"{TEST_EMPRESA_PREFIX}ORDEN_"
    )]
    test_names = [e["nombre"] for e in test_rows]
    assert test_names == [name_a, name_b, name_c]


@pytest.mark.asyncio
async def test_listar_empresas_ep_trimestre_decorativo(client):
    """V25 Cambio C (Capa 3): el `{trimestre}` del path es decorativo.
    El filtro pasa por `empresa.puedeSerEP` (ficha estructural), no por CT.

    Esto cambia la semántica del test anterior `_trimestre_inexistente`:
    antes esperábamos lista vacía con un trimestre bogus; ahora el filtro
    no depende del trimestre, así que dos trimestres distintos devuelven
    la MISMA lista. La promesa que sigue intacta: 200 OK (NO 404) y el
    trimestre del request se ecoa en la respuesta tal cual."""
    bogus_tri = "TEST-INEXISTENTE-9999"
    real_tri = TEST_TRIMESTRE

    resp_bogus = await client.get(
        f"/api/config-trimestral/{bogus_tri}/empresas-ep"
    )
    resp_real = await client.get(
        f"/api/config-trimestral/{real_tri}/empresas-ep"
    )
    assert resp_bogus.status_code == 200, resp_bogus.text
    assert resp_real.status_code == 200, resp_real.text

    # El trimestre se ecoa.
    assert resp_bogus.json()["trimestre"] == bogus_tri
    assert resp_real.json()["trimestre"] == real_tri

    # El filtro NO depende del trimestre — misma lista en ambos.
    ids_bogus = sorted(e["id"] for e in resp_bogus.json()["empresas"])
    ids_real = sorted(e["id"] for e in resp_real.json()["empresas"])
    assert ids_bogus == ids_real


@pytest.mark.asyncio
async def test_listar_empresas_ep_q2_real(client, db_session):
    """Smoke check on real data: should include empresas con `puedeSerEP=true`.

    V25 Cambio C (Capa 3): el gate migró a `empresa.puedeSerEP`. La migration
    0006 hizo backfill (empresa con CT.escuelaPropia=true en algún histórico
    → puedeSerEP=true), así que las IDs históricamente EP en 2026-Q2 deberían
    seguir presentes vía el nuevo filtro estructural.

    Skips when those IDs aren't present in this DB (e.g. local dev with
    different seed data) so the test stays portable.
    """
    expected_ids = {42, 49, 55, 60, 78, 82}

    # Confirm the seeded fixtures are actually present before asserting.
    # V25: chequeamos directamente el flag estructural en empresa.
    pre = await db_session.execute(
        text(
            'SELECT id FROM empresa '
            'WHERE "puedeSerEP" = true AND activa = true'
        ),
    )
    real_ids = {row["id"] for row in pre.mappings().all()}
    if not expected_ids.issubset(real_ids):
        pytest.skip(
            f"2026-Q2 EP fixtures not present in this DB "
            f"(found {sorted(real_ids)}); skipping real-data check."
        )

    resp = await client.get("/api/config-trimestral/2026-Q2/empresas-ep")
    assert resp.status_code == 200, resp.text
    data = resp.json()

    returned_ids = {e["id"] for e in data["empresas"]}
    assert expected_ids.issubset(returned_ids), (
        f"Expected {expected_ids} ⊆ returned, got {returned_ids}"
    )

    # Alphabetical order invariant on the real payload.
    nombres = [e["nombre"] for e in data["empresas"]]
    assert nombres == sorted(nombres)


# ── V25 Cambio C (Capa 3): /empresas-doble ─────────────────────


@pytest.mark.asyncio
async def test_listar_empresas_doble_devuelve_solo_marcadas(client, db_session):
    """GET /empresas-doble filtra por empresa.puedeSerDoble=true + activa=true.
    3 empresas: 2 con flag true, 1 sin flag → endpoint devuelve las 2 marcadas."""
    a_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}DOBLE_V25_A", activa=True,
    )
    b_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}DOBLE_V25_B", activa=True,
    )
    c_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}DOBLE_V25_C", activa=True,
    )
    await _set_puede_ser_doble(db_session, a_id)
    await _set_puede_ser_doble(db_session, b_id)
    # c_id keeps default puedeSerDoble=false.

    resp = await client.get(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-doble"
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()

    returned_ids = [e["id"] for e in data["empresas"]]
    assert a_id in returned_ids
    assert b_id in returned_ids
    assert c_id not in returned_ids
    assert data["total"] == len(data["empresas"])


@pytest.mark.asyncio
async def test_listar_empresas_doble_empresa_inactiva_excluida(client, db_session):
    """Empresa inactiva con puedeSerDoble=true → no aparece en el listado."""
    inactiva_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}DOBLE_V25_INACTIVA", activa=False,
    )
    await _set_puede_ser_doble(db_session, inactiva_id)

    resp = await client.get(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-doble"
    )
    assert resp.status_code == 200, resp.text
    returned_ids = [e["id"] for e in resp.json()["empresas"]]
    assert inactiva_id not in returned_ids


@pytest.mark.asyncio
async def test_listar_empresas_doble_sin_marcadas_lista_vacia(client, db_session):
    """Cuando ninguna empresa del test tiene puedeSerDoble=true → no aparece
    ninguna del prefijo TEST_EMPRESA_PREFIX. (Otras empresas reales de la BD
    pueden aparecer si tienen el flag activo — el assert solo checa que las
    empresas test no estén.)"""
    eid = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}DOBLE_V25_NO_FLAG", activa=True,
    )
    # No `_set_puede_ser_doble` — queda con default false.

    resp = await client.get(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-doble"
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert eid not in [e["id"] for e in data["empresas"]]


# ── V22 (Cambio A): permiteExtras ──────────────────────────────


@pytest.mark.asyncio
async def test_listar_empresas_permite_extras_filtra_correctamente(client, db_session):
    """GET /empresas-permite-extras returns only permiteExtras=true + activa=true.

    Three empresas: PE+activa, PE+inactiva, noPE+activa → only the first appears.
    """
    pe_activa_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_OK", activa=True,
    )
    pe_inactiva_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_INACTIVA", activa=False,
    )
    no_pe_id = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_NOPE", activa=True,
    )
    await _set_config_trimestral(
        db_session, pe_activa_id, escuela_propia=False, permite_extras=True,
    )
    await _set_config_trimestral(
        db_session, pe_inactiva_id, escuela_propia=False, permite_extras=True,
    )
    await _set_config_trimestral(
        db_session, no_pe_id, escuela_propia=False, permite_extras=False,
    )

    resp = await client.get(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/empresas-permite-extras",
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()

    returned_ids = [e["id"] for e in data["empresas"]]
    assert pe_activa_id in returned_ids
    assert pe_inactiva_id not in returned_ids
    assert no_pe_id not in returned_ids

    # The new endpoint is independent of escuelaPropia (we set ep=False above).
    row = next(e for e in data["empresas"] if e["id"] == pe_activa_id)
    assert row["nombre"] == f"{TEST_EMPRESA_PREFIX}PE_V22_OK"
    assert row["activa"] is True
    assert data["total"] == len(data["empresas"])


@pytest.mark.asyncio
async def test_listar_empresas_permite_extras_trimestre_inexistente(client):
    """Bogus trimestre → 200 with empty list."""
    resp = await client.get(
        "/api/config-trimestral/TEST-INEXISTENTE-V22-PE/empresas-permite-extras",
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["total"] == 0
    assert data["empresas"] == []


@pytest.mark.asyncio
async def test_get_trimestre_expone_permite_extras(client, db_session):
    """GET /{trimestre} surfaces the new permite_extras field per config row."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_GET")
    await _set_config_trimestral(
        db_session, eid, escuela_propia=False, permite_extras=True,
    )

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}")
    assert resp.status_code == 200, resp.text
    data = resp.json()

    row = next(c for c in data["configs"] if c["empresa_id"] == eid)
    assert row["permite_extras"] is True
    assert row["escuela_propia"] is False  # ortogonal


@pytest.mark.asyncio
async def test_resumen_incluye_contador_permite_extras(client, db_session):
    """GET /{trimestre}/resumen exposes permite_extras count."""
    a = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_R_A")
    b = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_R_B")
    c = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_R_C")
    await _set_config_trimestral(db_session, a, escuela_propia=False, permite_extras=True)
    await _set_config_trimestral(db_session, b, escuela_propia=False, permite_extras=True)
    await _set_config_trimestral(db_session, c, escuela_propia=False, permite_extras=False)

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}/resumen")
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert "permite_extras" in data
    # At least our 2 fixtures with permite_extras=true should be in the count
    # (other tests may have added more, so use >=).
    assert data["permite_extras"] >= 2


@pytest.mark.asyncio
async def test_batch_update_permite_extras(client, db_session):
    """PUT /batch toggles permiteExtras."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_BATCH")
    await _set_config_trimestral(
        db_session, eid, escuela_propia=False, permite_extras=False,
    )

    resp = await client.put(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/batch",
        json={"updates": [{"empresa_id": eid, "permite_extras": True}]},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["updated"] == 1

    # DB cross-check.
    row = await db_session.execute(
        text(
            'SELECT "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    assert row.scalar() is True


@pytest.mark.asyncio
async def test_put_individual_permite_extras(client, db_session):
    """PUT /{trimestre}/{empresa_id} toggles permiteExtras."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_PUT")
    await _set_config_trimestral(
        db_session, eid, escuela_propia=False, permite_extras=False,
    )

    resp = await client.put(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/{eid}",
        json={"permite_extras": True},
    )
    assert resp.status_code == 200, resp.text
    config = resp.json()["config"]
    assert config["permite_extras"] is True

    # DB cross-check.
    row = await db_session.execute(
        text(
            'SELECT "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    assert row.scalar() is True


@pytest.mark.asyncio
async def test_inicializar_defaults_hereda_acepta_extras(client, db_session):
    """POST /inicializar (modo defaults) → CT.permiteExtras = empresa.aceptaExtras.

    Uses a per-test trimestre so the test empresas have no pre-existing CT
    (inicializar only acts on empresas without a CT for the target trimestre).
    """
    new_tri = "TEST-V22-INIT-DEF"

    # Empresa A: aceptaExtras=true → expect CT.permiteExtras=true.
    a = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_INIT_A", acepta_extras=True,
    )
    # Empresa B: aceptaExtras=false → expect CT.permiteExtras=false.
    b = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_INIT_B", acepta_extras=False,
    )

    # Pre-clean the trimestre for these two empresas so inicializar acts.
    await db_session.execute(
        text(
            'DELETE FROM "configTrimestral" '
            'WHERE trimestre = :tri AND "empresaId" IN (:a, :b)'
        ),
        {"tri": new_tri, "a": a, "b": b},
    )
    await db_session.commit()

    resp = await client.post(
        f"/api/config-trimestral/{new_tri}/inicializar",
        json={},
    )
    assert resp.status_code == 200, resp.text

    row_a = await db_session.execute(
        text(
            'SELECT "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": a, "tri": new_tri},
    )
    row_b = await db_session.execute(
        text(
            'SELECT "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": b, "tri": new_tri},
    )
    assert row_a.scalar() is True
    assert row_b.scalar() is False

    # Clean up the spawned CTs so the trimestre stays disposable.
    await db_session.execute(
        text('DELETE FROM "configTrimestral" WHERE trimestre = :tri'),
        {"tri": new_tri},
    )
    await db_session.commit()


@pytest.mark.asyncio
async def test_inicializar_clonar_copia_permite_extras(client, db_session):
    """POST /inicializar with origen_trimestre clones permite_extras 1:1."""
    origen_tri = "TEST-V22-INIT-ORIG"
    destino_tri = "TEST-V22-INIT-DEST"

    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}PE_V22_CLONE")
    # Seed origen with permite_extras=true.
    await _set_config_trimestral(
        db_session, eid, escuela_propia=False, permite_extras=True,
        trimestre=origen_tri,
    )

    # Pre-clean destino just in case.
    await db_session.execute(
        text('DELETE FROM "configTrimestral" WHERE trimestre = :tri'),
        {"tri": destino_tri},
    )
    await db_session.commit()

    resp = await client.post(
        f"/api/config-trimestral/{destino_tri}/inicializar",
        json={"origen_trimestre": origen_tri},
    )
    assert resp.status_code == 200, resp.text

    row = await db_session.execute(
        text(
            'SELECT "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": destino_tri},
    )
    assert row.scalar() is True

    # Cleanup.
    await db_session.execute(
        text(
            'DELETE FROM "configTrimestral" WHERE trimestre IN (:o, :d)'
        ),
        {"o": origen_tri, "d": destino_tri},
    )
    await db_session.commit()


# ── V24 (Cambio B): frecuenciaEF / frecuenciaIT en endpoints CT ────


async def _set_ct_with_freq(
    db,
    empresa_id: int,
    *,
    freq_ef: int | None,
    freq_it: int | None,
    trimestre: str = TEST_TRIMESTRE,
):
    """Upsert CT with explicit freqEF/freqIT (and the V22 flags neutral)."""
    await db.execute(
        text(
            'INSERT INTO "configTrimestral" '
            '("empresaId", trimestre, "tipoParticipacion", "escuelaPropia", '
            '"permiteExtras", "frecuenciaEF", "frecuenciaIT", '
            '"disponibilidadDias", "updatedAt") '
            "VALUES (:eid, :tri, 'AMBAS', false, false, :ef, :it, 'L,M,X,J,V', NOW()) "
            'ON CONFLICT ("empresaId", trimestre) DO UPDATE SET '
            '    "frecuenciaEF" = EXCLUDED."frecuenciaEF", '
            '    "frecuenciaIT" = EXCLUDED."frecuenciaIT"'
        ),
        {"eid": empresa_id, "tri": trimestre, "ef": freq_ef, "it": freq_it},
    )
    await db.commit()


@pytest.mark.asyncio
async def test_get_trimestre_returns_freq_ef_it(client, db_session):
    """V24 Cambio B (T4): GET /{trimestre} expone frecuencia_ef y frecuencia_it
    para cada config en el JSON.
    """
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_GET")
    await _set_ct_with_freq(db_session, eid, freq_ef=4, freq_it=2)

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}")
    assert resp.status_code == 200, resp.text
    configs = resp.json()["configs"]
    target = next(c for c in configs if c["empresa_id"] == eid)
    assert target["frecuencia_ef"] == 4
    assert target["frecuencia_it"] == 2


@pytest.mark.asyncio
async def test_put_empresa_updates_freq_ef_it(client, db_session):
    """V24 Cambio B (T5): PUT /{trimestre}/{empresa_id} persiste
    frecuencia_ef y frecuencia_it cuando vienen en el body.
    """
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_PUT")
    await _set_ct_with_freq(db_session, eid, freq_ef=None, freq_it=None)

    resp = await client.put(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/{eid}",
        json={"frecuencia_ef": 5, "frecuencia_it": 3},
    )
    assert resp.status_code == 200, resp.text
    config = resp.json()["config"]
    assert config["frecuencia_ef"] == 5
    assert config["frecuencia_it"] == 3

    # DB cross-check
    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] == 5
    assert rec["frecuenciaIT"] == 3


@pytest.mark.asyncio
async def test_batch_update_freq_ef_it(client, db_session):
    """V24 Cambio B (T6): PUT /batch acepta frecuencia_ef y frecuencia_it.
    Toggle via batch para una empresa, DB cross-check.
    """
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_BATCH")
    await _set_ct_with_freq(db_session, eid, freq_ef=1, freq_it=0)

    resp = await client.put(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/batch",
        json={
            "updates": [
                {"empresa_id": eid, "frecuencia_ef": 6, "frecuencia_it": 4},
            ]
        },
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["updated"] == 1

    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] == 6
    assert rec["frecuenciaIT"] == 4


@pytest.mark.asyncio
async def test_inicializar_clonar_copies_freq_ef_it(client, db_session):
    """V24 Cambio B (T7): POST /inicializar modo clonar copia freqEF/freqIT 1:1
    desde el trimestre origen.
    """
    origen_tri = "TEST-V24-CLONE-ORIG-EFIT"
    destino_tri = "TEST-V24-CLONE-DEST-EFIT"
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_CLONE_EFIT")

    # Seed origen with freq_ef=7, freq_it=2.
    await _set_ct_with_freq(
        db_session, eid, freq_ef=7, freq_it=2, trimestre=origen_tri,
    )
    # Pre-clean destino.
    await db_session.execute(
        text('DELETE FROM "configTrimestral" WHERE trimestre = :tri'),
        {"tri": destino_tri},
    )
    await db_session.commit()

    resp = await client.post(
        f"/api/config-trimestral/{destino_tri}/inicializar",
        json={"origen_trimestre": origen_tri},
    )
    assert resp.status_code == 200, resp.text

    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": destino_tri},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] == 7
    assert rec["frecuenciaIT"] == 2

    # Cleanup the bespoke trimestres.
    await db_session.execute(
        text('DELETE FROM "configTrimestral" WHERE trimestre IN (:o, :d)'),
        {"o": origen_tri, "d": destino_tri},
    )
    await db_session.commit()


@pytest.mark.asyncio
async def test_inicializar_defaults_leaves_freq_ef_it_null(client, db_session):
    """V24 Cambio B (T7.2): POST /inicializar modo defaults NO rellena
    frecuenciaEF/IT (quedan NULL para que la planificadora los meta a mano).
    """
    new_tri = "TEST-V24-INIT-DEFAULTS-EFIT"
    eid = await _create_empresa(
        db_session, f"{TEST_EMPRESA_PREFIX}V24_INIT_DEFAULTS_EFIT",
    )

    # Asegurarse de que la empresa NO tiene CT previa en new_tri.
    await db_session.execute(
        text(
            'DELETE FROM "configTrimestral" WHERE trimestre = :tri AND "empresaId" = :eid'
        ),
        {"tri": new_tri, "eid": eid},
    )
    await db_session.commit()

    resp = await client.post(
        f"/api/config-trimestral/{new_tri}/inicializar",
        json={},
    )
    assert resp.status_code == 200, resp.text

    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": new_tri},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] is None
    assert rec["frecuenciaIT"] is None

    await db_session.execute(
        text('DELETE FROM "configTrimestral" WHERE trimestre = :tri'),
        {"tri": new_tri},
    )
    await db_session.commit()


# ── V24 (Cambio B, B4.5): bulk CT importer formato 10-col ─────────


def _build_v24_ct_excel(rows: list[dict], include_optional: bool = True) -> bytes:
    """Build a 10-col CT Excel for the V24 importer.
    `rows` items: {empresa, freq_ef, freq_it, escuela_propia, permite_extras,
                   tipo?, dias?, turno?, voluntarios?, notas?}
    `include_optional=False` → solo las 5 columnas requeridas (sirve para
    verificar que el importer acepta el mínimo).
    """
    import openpyxl
    from io import BytesIO

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "ConfigTrimestral"

    if include_optional:
        headers = [
            "Empresa", "Frecuencia EF", "Frecuencia IT", "Tipo", "Dias",
            "Turno", "Voluntarios", "Escuela Propia", "Permite Extras", "Notas",
        ]
    else:
        headers = [
            "Empresa", "Frecuencia EF", "Frecuencia IT",
            "Escuela Propia", "Permite Extras",
        ]
    ws.append(headers)

    for r in rows:
        if include_optional:
            ws.append([
                r["empresa"],
                r.get("freq_ef"),
                r.get("freq_it"),
                r.get("tipo"),
                r.get("dias"),
                r.get("turno"),
                r.get("voluntarios"),
                "SI" if r.get("escuela_propia") else "NO",
                "SI" if r.get("permite_extras") else "NO",
                r.get("notas"),
            ])
        else:
            ws.append([
                r["empresa"],
                r.get("freq_ef"),
                r.get("freq_it"),
                "SI" if r.get("escuela_propia") else "NO",
                "SI" if r.get("permite_extras") else "NO",
            ])

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


@pytest.mark.asyncio
async def test_importar_excel_ct_format_nuevo(client, db_session):
    """V24 B4.5: el importer acepta el nuevo formato 10-columnas y devuelve
    `formato_detectado='v24-split-efit'`. Verifica dry_run + apply en el
    mismo test contra un CT pre-existente (UPDATE path)."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_IMP_UPDATE")
    # Sembrar un CT en TEST_TRIMESTRE para forzar la rama UPDATE.
    await _set_ct_with_freq(db_session, eid, freq_ef=None, freq_it=None)

    excel = _build_v24_ct_excel([{
        "empresa": f"{TEST_EMPRESA_PREFIX}V24_IMP_UPDATE",
        "freq_ef": 4, "freq_it": 2, "tipo": "AMBAS",
        "dias": "L,M,X,J,V", "turno": "M", "voluntarios": 1,
        "escuela_propia": False, "permite_extras": True, "notas": "Test V24",
    }])

    # dry_run=True → preview, no aplica.
    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "ct_v24.xlsx", excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "true"},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["formato_detectado"] == "v24-split-efit"
    assert data["total_procesados"] == 1
    assert data["aplicados"] == 0
    assert data["dry_run"] is True

    target = next(p for p in data["preview"] if p["empresa_id"] == eid)
    assert target["frecuencia_ef"] == 4
    assert target["frecuencia_it"] == 2
    assert target["permite_extras"] is True
    assert target["escuela_propia"] is False

    # BD inalterada por el dry-run.
    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT", "permiteExtras" '
            'FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] is None  # sigue NULL por _set_ct_with_freq inicial
    assert rec["frecuenciaIT"] is None
    assert rec["permiteExtras"] is False

    # dry_run=False → aplica.
    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "ct_v24.xlsx", excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "false"},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()
    assert data["aplicados"] == 1

    # BD post-aplicación: valores persistidos.
    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT", "permiteExtras", "escuelaPropia" '
            'FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] == 4
    assert rec["frecuenciaIT"] == 2
    assert rec["permiteExtras"] is True
    assert rec["escuelaPropia"] is False


@pytest.mark.asyncio
async def test_importar_excel_ep_pe_override_total(client, db_session):
    """V24 B4.5 (D3): EP y PE override total — Excel manda incluso si el CT
    existente ya tenía un valor distinto. Pasar SI/NO sobreescribe."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_IMP_OVERRIDE")
    # CT inicial con EP=true, PE=true.
    await db_session.execute(
        text(
            'INSERT INTO "configTrimestral" '
            '("empresaId", trimestre, "tipoParticipacion", "escuelaPropia", '
            '"permiteExtras", "disponibilidadDias", "updatedAt") '
            "VALUES (:eid, :tri, 'AMBAS', true, true, 'L,M,X,J,V', NOW()) "
            'ON CONFLICT ("empresaId", trimestre) DO UPDATE SET '
            '    "escuelaPropia" = true, "permiteExtras" = true'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    await db_session.commit()

    # Excel manda EP=NO, PE=NO → debe ganar.
    excel = _build_v24_ct_excel([{
        "empresa": f"{TEST_EMPRESA_PREFIX}V24_IMP_OVERRIDE",
        "freq_ef": 1, "freq_it": 0,
        "escuela_propia": False, "permite_extras": False,
    }])
    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "ct_v24.xlsx", excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "false"},
    )
    assert resp.status_code == 200, resp.text

    row = await db_session.execute(
        text(
            'SELECT "escuelaPropia", "permiteExtras" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["escuelaPropia"] is False, "Excel debió sobreescribir EP=true → false"
    assert rec["permiteExtras"] is False, "Excel debió sobreescribir PE=true → false"


@pytest.mark.asyncio
async def test_importar_excel_freq_null_se_persiste_como_null(client, db_session):
    """V24 B4.5: si una celda Freq EF / Freq IT viene vacía → NULL en BD
    (D2 + D7: NULL significa 'no participa en ese tipo')."""
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_IMP_NULL")
    await _set_ct_with_freq(db_session, eid, freq_ef=99, freq_it=99)  # ruido previo

    excel = _build_v24_ct_excel([{
        "empresa": f"{TEST_EMPRESA_PREFIX}V24_IMP_NULL",
        "freq_ef": None, "freq_it": 3,  # EF vacío, IT=3
        "escuela_propia": False, "permite_extras": False,
    }])
    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "ct_v24.xlsx", excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "false"},
    )
    assert resp.status_code == 200, resp.text

    row = await db_session.execute(
        text(
            'SELECT "frecuenciaEF", "frecuenciaIT" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    rec = row.mappings().first()
    assert rec["frecuenciaEF"] is None
    assert rec["frecuenciaIT"] == 3


@pytest.mark.asyncio
async def test_importar_excel_rechaza_formato_viejo(client):
    """V24 B4.5 (D2): el formato legacy/ideal anterior (columna 'Frecuencia'
    única, sin EF/IT split, sin EP/PE) debe ser rechazado con 400 y mensaje
    claro sobre las columnas requeridas."""
    import openpyxl
    from io import BytesIO

    # Excel formato viejo "ideal": Empresa | Frecuencia | Tipo | Dias | ...
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Empresa", "Frecuencia", "Tipo", "Dias", "Turno", "Voluntarios", "Notas"])
    ws.append(["ACME", 5, "AMBAS", "L,M,X,J,V", "M", 1, "legacy row"])
    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "legacy.xlsx", buf.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "true"},
    )
    assert resp.status_code == 400, resp.text
    detail = resp.json()["detail"].lower()
    assert "freq_ef" in detail or "frecuencia ef" in detail or "requeridas" in detail


@pytest.mark.asyncio
async def test_export_excel_includes_freq_ef_it_ep_pe_columns(client, db_session):
    """V24 Cambio B (T8, decisión D9): GET /exportar-excel devuelve un xlsx
    con la nueva fila de headers (10 cols), sin "Frecuencia" single, con
    "Frecuencia EF", "Frecuencia IT", "Escuela Propia", "Permite Extras".
    """
    from io import BytesIO
    import openpyxl

    # Sembrar al menos 1 fila para que el export tenga data.
    eid = await _create_empresa(db_session, f"{TEST_EMPRESA_PREFIX}V24_EXPORT")
    await _set_ct_with_freq(db_session, eid, freq_ef=4, freq_it=2)

    resp = await client.get(f"/api/config-trimestral/{TEST_TRIMESTRE}/exportar-excel")
    assert resp.status_code == 200, resp.text

    wb = openpyxl.load_workbook(BytesIO(resp.content), data_only=True)
    ws = wb.active
    headers = [cell.value for cell in ws[1]]

    assert headers == [
        "Empresa", "Frecuencia EF", "Frecuencia IT", "Tipo", "Dias", "Turno",
        "Voluntarios", "Escuela Propia", "Permite Extras", "Notas",
    ]
    assert "Frecuencia" not in headers, "la columna single 'Frecuencia' fue eliminada en D9"

    # Verifica que la fila de la empresa de test trae los valores correctos.
    found = False
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[0] == f"{TEST_EMPRESA_PREFIX}V24_EXPORT":
            assert row[1] == 4   # Frecuencia EF
            assert row[2] == 2   # Frecuencia IT
            assert row[7] == "NO"  # Escuela Propia (false → NO)
            assert row[8] == "NO"  # Permite Extras (false → NO)
            found = True
            break
    assert found, "la empresa de test no aparece en el xlsx"


# ─────────────────────────────────────────────────────────────
# V25 Capa 9 — Guard explícito activa/inactiva (decisión C2)
# ─────────────────────────────────────────────────────────────


@pytest.mark.asyncio
async def test_put_ct_rechaza_empresa_inactiva(client, db_session):
    """V25 Capa 9 (decisión C2): PUT /{tri}/{empresa_id} contra empresa
    inactiva → 422 con el nombre en el mensaje. Antes pasaba silenciosamente
    (solo era filtrado por el solver y /calcular). Ahora falla temprano.
    """
    nombre = f"{TEST_EMPRESA_PREFIX}V25_CT_INACTIVA"
    eid = await _create_empresa(db_session, nombre, activa=False)

    resp = await client.put(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/{eid}",
        json={"escuela_propia": True},
    )
    assert resp.status_code == 422, resp.text
    assert nombre in resp.json()["detail"], (
        f"el detail debería incluir el nombre de la empresa: {resp.json()}"
    )

    # BD cross-check: CT NO se creó.
    row = await db_session.execute(
        text(
            'SELECT id FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid, "tri": TEST_TRIMESTRE},
    )
    assert row.first() is None, "CT no debería existir para empresa inactiva"


@pytest.mark.asyncio
async def test_bulk_ct_importar_excel_skip_empresa_inactiva(client, db_session):
    """V25 Capa 9: bulk /importar-excel rechaza por fila la empresa inactiva
    con warning específico ('inactiva en BD'), las demás filas se procesan.
    Batch NO aborta.
    """
    activa = f"{TEST_EMPRESA_PREFIX}V25_CT_BULK_ACTIVA"
    inactiva = f"{TEST_EMPRESA_PREFIX}V25_CT_BULK_INACTIVA"
    eid_a = await _create_empresa(db_session, activa, activa=True)
    eid_i = await _create_empresa(db_session, inactiva, activa=False)

    excel = _build_v24_ct_excel([
        {
            "empresa": activa,
            "freq_ef": 3, "freq_it": 1, "tipo": "AMBAS",
            "dias": "L,M,X,J,V", "turno": "M", "voluntarios": 1,
            "escuela_propia": False, "permite_extras": False, "notas": None,
        },
        {
            "empresa": inactiva,
            "freq_ef": 2, "freq_it": 0, "tipo": "EF",
            "dias": "L,M,X,J,V", "turno": "M", "voluntarios": 1,
            "escuela_propia": False, "permite_extras": False, "notas": None,
        },
    ])

    resp = await client.post(
        f"/api/config-trimestral/{TEST_TRIMESTRE}/importar-excel",
        files={
            "file": (
                "ct_v25_capa9.xlsx", excel,
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        },
        data={"dry_run": "false"},
    )
    assert resp.status_code == 200, resp.text
    data = resp.json()

    # La activa sí se procesa, la inactiva NO.
    assert data["aplicados"] == 1, (
        f"esperaba 1 aplicado (solo la activa), hay {data['aplicados']}"
    )
    warnings_inactiva = [
        w for w in data["warnings"] if inactiva in w and "inactiva" in w.lower()
    ]
    assert warnings_inactiva, (
        f"warning específico para '{inactiva}' (inactiva) no encontrado: "
        f"{data['warnings']}"
    )

    # BD cross-check.
    row_a = await db_session.execute(
        text(
            'SELECT "frecuenciaEF" FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid_a, "tri": TEST_TRIMESTRE},
    )
    assert row_a.scalar() == 3, "la activa debería tener freq_ef=3"

    row_i = await db_session.execute(
        text(
            'SELECT id FROM "configTrimestral" '
            'WHERE "empresaId" = :eid AND trimestre = :tri'
        ),
        {"eid": eid_i, "tri": TEST_TRIMESTRE},
    )
    assert row_i.first() is None, "CT no debería existir para empresa inactiva"
